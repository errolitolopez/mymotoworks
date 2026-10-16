"""Reading the old spreadsheet and the import-items command."""

import os
import zipfile

import click
from flask.cli import with_appcontext
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from mymotoworks.database import get_db
from mymotoworks.errors import ValidationError
from mymotoworks.services import ItemImportService

SHEET_NAME = "STOCK INVENTORY"
# The heading text in the sheet for each thing we read. REMAINING STOCK, not STOCK: it is the
# sheet's own total after sales. The RE-ORDER column is not read at all.
HEADINGS = {"name": "ITEM NAME", "model": "MODEL", "part_no": "PART NO.", "puhunan": "PUHUNAN",
            "benta": "BENTA", "stock": "REMAINING STOCK"}
MAX_HEADING_SEARCH_ROWS = 10


def read_inventory_rows(path):
    """The item rows of the STOCK INVENTORY sheet as (row number, record) pairs, where the record
    has the keys name, model, part_no, puhunan, benta, stock. Raises ValidationError with a clear
    message when the file or sheet is not as expected."""
    if not os.path.isfile(path):
        raise ValidationError(f"Cannot find the file: {path}")
    try:
        # data_only: read the numbers Excel last calculated, not the formulas.
        workbook = load_workbook(path, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError):
        raise ValidationError("Cannot read that file. Is it an Excel .xlsx file?") from None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValidationError(f'The file has no sheet named "{SHEET_NAME}".')
        sheet_rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()

    for index, row in enumerate(sheet_rows[:MAX_HEADING_SEARCH_ROWS]):
        headings = [str(cell).strip().upper() if cell is not None else "" for cell in row]
        if HEADINGS["name"] in headings:
            break
    else:
        raise ValidationError(f'Cannot find the "{HEADINGS["name"]}" heading row in sheet "{SHEET_NAME}".')
    columns = {}
    for key, text in HEADINGS.items():
        if text not in headings:
            raise ValidationError(f'The sheet has no "{text}" column.')
        columns[key] = headings.index(text)

    rows = []
    for number, row in enumerate(sheet_rows[index + 1:], start=index + 2):
        rows.append((number, {key: row[column] if column < len(row) else None
                              for key, column in columns.items()}))
    return rows


def _money(centavos):
    return f"{centavos / 100:,.2f}"


def describe_plan(plan, verbose=False):
    """The preview the owner reads before anything is saved."""
    lines = [f'Read {plan.rows_read} rows from "{SHEET_NAME}".']
    count = len(plan.items)
    lines.append(f"Will import {count} item{'' if count == 1 else 's'}.")
    if plan.problems:
        lines.append(f"Skipped {len(plan.problems)} rows (add these by hand on the Items page):")
        lines += [f"  Row {number}: {reason}" for number, reason in plan.problems]
    notes = list(plan.warnings)
    if plan.no_puhunan:
        notes.insert(0, f"{plan.no_puhunan} item{' has' if plan.no_puhunan == 1 else 's have'} "
                        "no puhunan (counted as 0).")
    if notes:
        lines.append("Notes:")
        lines += [f"  {note}" for note in notes]
    lines.append(f'Stock is the "{HEADINGS["stock"]}" column. Re-order is 0 for every item.')
    if verbose:
        lines.append("Items:")
        for number, fields in plan.items:
            lines.append(
                f"  Row {number}: {fields['name']} | {fields['model']} | {fields['part_no']} | "
                f"puhunan {_money(fields['puhunan'])} | benta {_money(fields['benta'])} | "
                f"stock {fields['stock']} | re-order {fields['reorder_level']}")
    if plan.existing:
        lines.append(f"The shop already has {plan.existing} item{'' if plan.existing == 1 else 's'}, "
                     "so the real import would be refused.")
    return "\n".join(lines)


@click.command("import-items")
@click.argument("path", type=click.Path())
@click.option("--dry-run", is_flag=True, help="Only show what would be imported.")
@click.option("-v", "--verbose", is_flag=True, help="Also list every item that would be imported.")
@click.option("--yes", is_flag=True, help="Do not ask before importing.")
@with_appcontext
def import_items_command(path, dry_run, verbose, yes):
    """Load the items and current stock from the old spreadsheet (PATH, an .xlsx file)."""
    service = ItemImportService(get_db())
    try:
        plan = service.plan(read_inventory_rows(path))
    except ValidationError as error:
        raise click.ClickException(error.message) from None
    click.echo(describe_plan(plan, verbose))
    if dry_run:
        return
    try:
        if not plan.items:
            raise ValidationError("There is nothing to import.")
        if not yes and not click.confirm(f"Import {len(plan.items)} items?"):
            click.echo("Nothing imported.")
            return
        added = service.run(plan)
    except ValidationError as error:
        raise click.ClickException(error.message) from None
    click.echo(f"Imported {added} items.")
