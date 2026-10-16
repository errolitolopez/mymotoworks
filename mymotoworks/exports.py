"""Excel files for the reports. Money cells are numbers (pesos), and text is never turned into a formula."""

from io import BytesIO

from openpyxl import Workbook

MONEY_FORMAT = "#,##0.00"
# A text that starts with one of these would be read by Excel as a formula (or a command).
FORMULA_STARTS = ("=", "+", "-", "@")


def _pesos(centavos):
    """10350 -> 103.5 (a number, so Excel can add it up)."""
    return centavos / 100


def _write(sheet, values, money=()):
    """Add one row. Text goes in as text, never as a formula. `money` is the 0-based columns that hold centavos."""
    sheet.append([None] * len(values))
    row = sheet.max_row
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index)
        if isinstance(value, str):
            cell.value = value
            if value.startswith(FORMULA_STARTS):
                cell.data_type = "s"
        elif index - 1 in money and value is not None:
            cell.value = _pesos(value)
            cell.number_format = MONEY_FORMAT
        else:
            cell.value = value


def _widths(sheet, widths):
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width


def _bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def sales_workbook(report, months):
    """The sales report of a month (`ReportService.sales`) and the months of its year (`.months`)."""
    workbook = Workbook()
    days = workbook.active
    days.title = "Days"
    _write(days, ["Date", "Sales", "Pieces", "Total", "Cost", "Tubo"])
    for day in report["days"]:
        _write(days, [day["date"], day["sales"], day["pieces"], day["total"], day["cost"], day["tubo"]], money=(3, 4, 5))
    totals = report["totals"]
    _write(days, ["Total", totals["sales"], totals["pieces"], totals["total"], totals["cost"], totals["tubo"]],
           money=(3, 4, 5))
    days.append([])
    _write(days, ["Spent on restocks", None, None, report["restocked"]], money=(3,))
    _widths(days, {"A": 18, "B": 8, "C": 8, "D": 14, "E": 14, "F": 14})

    best = workbook.create_sheet("Best sellers")
    _write(best, ["Item", "Pieces", "Total", "Tubo"])
    for item in report["best_sellers"]:
        _write(best, [item["name"], item["pieces"], item["total"], item["tubo"]], money=(2, 3))
    _widths(best, {"A": 40, "B": 8, "C": 14, "D": 14})

    by_month = workbook.create_sheet("Months")
    _write(by_month, ["Month", "Sales", "Pieces", "Total", "Cost", "Tubo"])
    for entry in months["months"]:
        _write(by_month, [entry["month"], entry["sales"], entry["pieces"], entry["total"], entry["cost"], entry["tubo"]],
               money=(3, 4, 5))
    _widths(by_month, {"A": 12, "B": 8, "C": 8, "D": 14, "E": 14, "F": 14})
    return _bytes(workbook)


def stock_workbook(report, items):
    """The stock report (`ReportService.stock`) and every item that counts stock (`ReportService.stock_items`)."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    totals = report["totals"]
    _write(summary, ["Capital in stock", totals["capital"]], money=(1,))
    _write(summary, ["If all sold (revenue)", totals["revenue"]], money=(1,))
    _write(summary, ["Tubo if all sold", totals["tubo"]], money=(1,))
    _write(summary, ["Items", totals["items"]])
    _write(summary, ["Pieces", totals["pieces"]])
    _write(summary, ["Low stock", totals["low_stock"]])
    _write(summary, ["Out of stock", totals["out_of_stock"]])
    _write(summary, ["Below zero", totals["below_zero"]])
    _widths(summary, {"A": 26, "B": 16})

    low = workbook.create_sheet("Low stock")
    _write(low, ["Item", "Model", "Part no.", "Stock", "Re-order at"])
    for item in report["low_stock"]:
        _write(low, [item["name"], item["model"], item["part_no"], item["stock"], item["reorder_level"]])
    _widths(low, {"A": 40, "B": 18, "C": 14, "D": 8, "E": 12})

    every = workbook.create_sheet("All items")
    _write(every, ["Item", "Model", "Part no.", "Puhunan", "Benta", "Stock", "Capital", "Revenue if sold"])
    for item in items:
        stock = max(item["stock"], 0)
        _write(every, [item["name"], item["model"], item["part_no"], item["puhunan"], item["benta"], item["stock"],
                       item["puhunan"] * stock, item["benta"] * stock], money=(3, 4, 6, 7))
    _widths(every, {"A": 40, "B": 18, "C": 14, "D": 12, "E": 12, "F": 8, "G": 14, "H": 16})
    return _bytes(workbook)
